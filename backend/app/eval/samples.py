"""Generates the synthetic DOCX golden sample with an embedded swimlane flowchart.

    uv run python -m app.eval.samples

Covers what the other golden documents don't: the DOCX path, embedded images and a process
diagram that disagrees with the text (seeded gap XD004), with a dead end drawn in it (ST002).
"""

import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from app.eval.golden import GOLDEN_DIR

CASE_DIR = GOLDEN_DIR / "sample-access-request"
DOCX_NAME = "SOP-IT-021.docx"

LANES = ["Requester", "Line Manager", "IT Service Desk"]
# id, label, kind, lane index, x centre
SHAPES = [
    ("d1", "Start", "start", 0, 270),
    ("d2", "Submit access\nrequest form", "task", 0, 560),
    ("d3", "Approve?", "decision", 1, 860),
    ("d7", "Reject and notify\nrequester", "task", 1, 1220),
    ("d4", "Provision access", "task", 2, 860),
    ("d5", "Notify requester", "task", 2, 1180),
    ("d6", "End", "end", 2, 1460),
]
ARROWS = [
    ("d1", "d2", None),
    ("d2", "d3", None),
    ("d3", "d4", "Yes"),
    ("d3", "d7", "No"),
    ("d4", "d5", None),
    ("d5", "d6", None),
]  # d7 deliberately leads nowhere

W, H, LANE_H, LABEL_W = 1600, 780, 240, 150


def _font(size: int):
    for name in ("arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def draw_flowchart(path: Path) -> None:
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    font, small = _font(26), _font(22)
    top = 40
    d.text((W // 2, 18), "Figure 1: Access request flow", fill="black", font=font, anchor="mm")
    for i, lane in enumerate(LANES):
        y0 = top + i * LANE_H
        d.rectangle([10, y0, W - 10, y0 + LANE_H], outline="black", width=2)
        d.rectangle([10, y0, LABEL_W, y0 + LANE_H], outline="black", width=2, fill="#eeeeee")
        d.multiline_text(
            (LABEL_W // 2 + 5, y0 + LANE_H // 2),
            lane.replace(" ", "\n", 1),
            fill="black",
            font=small,
            anchor="mm",
            align="center",
        )
    centre = {}
    for sid, label, kind, lane, x in SHAPES:
        y = top + lane * LANE_H + LANE_H // 2
        centre[sid] = (x, y)
        if kind in ("start", "end"):
            d.ellipse([x - 80, y - 38, x + 80, y + 38], outline="black", width=3)
        elif kind == "decision":
            d.polygon(
                [(x, y - 70), (x + 110, y), (x, y + 70), (x - 110, y)], outline="black", width=3
            )
        else:
            d.rounded_rectangle(
                [x - 120, y - 50, x + 120, y + 50], radius=12, outline="black", width=3
            )
        d.multiline_text((x, y), label, fill="black", font=small, anchor="mm", align="center")
    for a, b, label in ARROWS:
        (x0, y0), (x1, y1) = centre[a], centre[b]
        if y0 == y1:
            sx, ex = (x0 + 120, x1 - 120) if x1 > x0 else (x0 - 120, x1 + 120)
            d.line([sx, y0, ex, y1], fill="black", width=3)
            d.polygon([(ex, y1), (ex - 16, y1 - 9), (ex - 16, y1 + 9)], fill="black")
            mid = ((sx + ex) // 2, y0 - 18)
        else:
            sy, ey = y0 + 70, y1 - 50
            d.line([x0, sy, x0, (sy + ey) // 2, x1, (sy + ey) // 2, x1, ey], fill="black", width=3)
            d.polygon([(x1, ey), (x1 - 9, ey - 16), (x1 + 9, ey - 16)], fill="black")
            mid = (x1 + 30, (sy + ey) // 2 - 18)
        if label:
            d.text(mid, label, fill="black", font=small, anchor="mm")
    img.save(path)


PARAGRAPHS = [
    ("heading1", "1 Purpose"),
    ("text", "1.1 This procedure describes how employees request access to company IT systems."),
    ("heading1", "2 Procedure"),
    ("text", "2.1 The requester submits the access request form in the IT portal."),
    ("text", "2.2 The line manager approves or rejects the request within 2 business days."),
    (
        "text",
        "2.3 For privileged (administrator) access, the Information Security team reviews "
        "the request after manager approval.",
    ),
    ("text", "2.4 The IT Service Desk provisions access and notifies the requester."),
    ("text", "2.5 Access must be reviewed periodically."),
    ("figure", "Figure 1: Access request flow"),
    ("heading1", "3 Approval matrix"),
    (
        "table",
        [
            ["Access type", "Approver"],
            ["Standard", "Line manager"],
            ["Privileged", "Line manager and Information Security"],
            ["Emergency", "IT Duty Manager"],
        ],
    ),
]

EXPECTED = [
    {
        "code": "XD004",
        "document": DOCX_NAME,
        "quote": "For privileged (administrator) access, the Information Security team reviews the "
        "request after manager approval.",
        "note": "The security review step is missing from Figure 1.",
    },
    {
        "code": "ST002",
        "document": DOCX_NAME,
        "quote": "The line manager approves or rejects the request within 2 business days.",
        "note": "After rejection nothing happens: in Figure 1 the reject step leads nowhere.",
    },
    {
        "code": "TM003",
        "document": DOCX_NAME,
        "quote": "Access must be reviewed periodically.",
        "note": "'Periodically' sets no interval.",
    },
    {
        "code": "RS001",
        "document": DOCX_NAME,
        "quote": "Access must be reviewed periodically.",
        "note": "No one owns the access review.",
    },
]


def build(case_dir: Path = CASE_DIR) -> Path:
    from docx import Document
    from docx.shared import Inches

    case_dir.mkdir(parents=True, exist_ok=True)
    chart = case_dir / "figure-1-access-request-flow.png"
    draw_flowchart(chart)

    doc = Document()
    doc.add_heading("SOP-IT-021: User Access Request", level=0)
    doc.add_paragraph(
        "SYNTHETIC SAMPLE for testing the policy compiler. Version 1.0. "
        "Owner: IT Service Management."
    )
    for kind, content in PARAGRAPHS:
        if kind == "heading1":
            doc.add_heading(content, level=1)
        elif kind == "text":
            doc.add_paragraph(content)
        elif kind == "figure":
            doc.add_picture(str(chart), width=Inches(6.3))
            doc.add_paragraph(content, style="Caption")
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    out = case_dir / DOCX_NAME
    doc.save(out)

    manifest = {
        "id": "sample-access-request",
        "title": "SOP-IT-021 User Access Request (DOCX with flowchart)",
        "synthetic": True,
        "documents": [DOCX_NAME],
        "expected_findings": EXPECTED,
    }
    (case_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return out


if __name__ == "__main__":
    print(build())
    sys.exit(0)
